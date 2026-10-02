# -*- coding: utf-8 -*-
"""Доводка docx до рабочего профиля СТП БГУИР 01-2024 (навык bsuire-report): символы, латиница курсивом, таблицы, перечисления, интервалы."""
import copy
import re

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

LATIN = re.compile(r"[A-Za-z](?:[A-Za-z0-9_+#]|[.\-](?=[A-Za-z0-9]))*")
CHUNK = re.compile(r"\S+|\s+")
UNITS = r"(?:%|USD|BYN|RUB|EUR|тыс\.|млн|млрд|мин|мс|ч|шт\.|мм|см|кг|Гб|МБ|Мб|GB|MB)"
NBSP = " "


def _is_heading(p):
    return (p.style.name or "").lower().startswith(("heading", "toc"))


def _has_field(r):
    x = r._r
    return x.find(qn("w:fldChar")) is not None or x.find(qn("w:instrText")) is not None or x.find(qn("w:drawing")) is not None


def clean_text(t):
    t = t.replace("ё", "е").replace("Ё", "Е").replace("—", "–")
    t = re.sub(r"(\d) (" + UNITS + r")(?![А-Яа-яA-Za-z])", lambda m: m.group(1) + NBSP + m.group(2), t)
    t = re.sub(r"(?<=\S)  +(?=\S)", " ", t)
    return t


def split_latin(text):
    """[(фрагмент, курсив)]: латиница курсивом, пути, URL и адреса не трогаются."""
    out = []
    for chunk in CHUNK.findall(text):
        if chunk.isspace() or any(c in chunk for c in ("/", "\\", "@", "://")):
            out.append((chunk, False))
            continue
        pos = 0
        for m in LATIN.finditer(chunk):
            if m.start() > pos:
                out.append((chunk[pos:m.start()], False))
            out.append((m.group(0), True))
            pos = m.end()
        if pos < len(chunk):
            out.append((chunk[pos:], False))
    merged = []
    for t, it in out:
        if merged and merged[-1][1] == it:
            merged[-1] = (merged[-1][0] + t, it)
        else:
            merged.append((t, it))
    return merged


def italic_latin(p):
    if _is_heading(p):
        return
    for r in list(p.runs):
        if _has_field(r) or not r.text:
            continue
        rs = (r.style.name or "") if r.style is not None else ""
        if "Verbatim" in rs or "Code" in rs:
            continue
        parts = split_latin(r.text)
        if len(parts) == 1:
            r.italic = True if parts[0][1] else r.italic
            continue
        prev = r._r
        for i, (t, it) in enumerate(parts):
            if i == 0:
                r.text = t
                r.italic = True if it else None
                continue
            new = copy.deepcopy(r._r)
            prev.addnext(new)
            from docx.text.run import Run
            nr = Run(new, p)
            nr.text = t
            nr.italic = True if it else None
            prev = new


def fix_chars(p):
    for r in p.runs:
        if r.text and not _has_field(r):
            nt = clean_text(r.text)
            if nt != r.text:
                r.text = nt


def table_profile(t):
    tbl_pr = t._tbl.tblPr
    for old in tbl_pr.findall(qn("w:jc")):
        tbl_pr.remove(old)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tbl_pr.append(jc)
    for i, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0 and tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.first_line_indent = Mm(0)
                pf.left_indent = Mm(0)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.keep_with_next = False
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                fix_chars(p)
                italic_latin(p)


def list_profile(p):
    ppr = p._p.pPr
    if ppr is None or ppr.find(qn("w:numPr")) is None:
        return False
    pf = p.paragraph_format
    pf.left_indent = Mm(12.5)
    pf.first_line_indent = Mm(-7.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return True


def _is_empty(p):
    x = p._p
    if p.text.strip():
        return False
    if x.findall(".//" + qn("w:drawing")) or x.find(".//" + qn("w:br")) is not None or x.find(".//" + qn("w:fldChar")) is not None or x.find(".//" + qn("w:instrText")) is not None:
        return False
    return x.pPr is None or x.pPr.find(qn("w:sectPr")) is None


def empties_to_spacing(doc):
    """Пустые абзацы титульного листа заменяются интервалом перед следующим абзацем (18 pt на абзац), прочие пустые абзацы удаляются."""
    pars = list(doc.paragraphs)
    toc = next((i for i, p in enumerate(pars) if p.text.strip().upper().startswith("СОДЕРЖАНИЕ")), len(pars))
    acc = 0
    for i, p in enumerate(pars):
        if _is_empty(p):
            if i < toc:
                acc += 18
            p._p.getparent().remove(p._p)
            continue
        if acc and i < toc:
            p.paragraph_format.space_before = Pt((p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0) + acc)
            acc = 0


def polish(doc):
    empties_to_spacing(doc)
    for t in doc.tables:
        table_profile(t)
    for p in doc.paragraphs:
        fix_chars(p)
        if not list_profile(p) and not _is_heading(p):
            pf = p.paragraph_format
            if pf.space_before in (Pt(12),) or pf.space_after in (Pt(12),):
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
        italic_latin(p)
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                fix_chars(p)
