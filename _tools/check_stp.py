# -*- coding: utf-8 -*-
"""Проверка docx по рабочему профилю СТП БГУИР 01-2024 (навык bsuire-report): PASS / FAIL / MANUAL по пунктам."""
import glob
import os
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LATIN = re.compile(r"[A-Za-z]")


def body_pars(doc):
    for p in doc.paragraphs:
        st = (p.style.name or "").lower()
        if st.startswith(("heading", "toc")) or not p.text.strip():
            continue
        yield p


def check(path):
    doc = Document(path)
    res = {}
    sec = doc.sections[0]
    mm = lambda v: round(v.mm)
    res["Страница и поля"] = ("PASS" if (mm(sec.page_width), mm(sec.page_height), mm(sec.left_margin), mm(sec.right_margin), mm(sec.top_margin), mm(sec.bottom_margin)) == (210, 297, 30, 15, 20, 20)
                              else f"FAIL {(mm(sec.page_width), mm(sec.page_height), mm(sec.left_margin), mm(sec.right_margin), mm(sec.top_margin), mm(sec.bottom_margin))}")
    # основной текст: берём абзацы, не заголовки, не списки, не подписи, не титул
    main = []
    started = False
    for p in doc.paragraphs:
        st = (p.style.name or "").lower()
        if st.startswith("heading") and p.text.strip().upper().startswith("1 "):
            started = True
        if not started or st.startswith(("heading", "toc")) or not p.text.strip():
            continue
        if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None:
            continue
        if re.match(r"^(Таблица|Рисунок) [\dА-Я.]+ [–-]", p.text) or p.alignment == 1:
            continue
        if len(p.text) < 60:
            continue
        main.append(p)
    bad = {"font": 0, "size": 0, "line": 0, "before": 0, "after": 0, "indent": 0, "align": 0}
    for p in main:
        pf = p.paragraph_format
        runs = [r for r in p.runs if r.text.strip()]
        if any((r.font.name or "Times New Roman") != "Times New Roman" for r in runs):
            bad["font"] += 1
        if any(r.font.size is not None and r.font.size != Pt(14) for r in runs):
            bad["size"] += 1
        if pf.line_spacing != Pt(18) or pf.line_spacing_rule != 4:
            bad["line"] += 1
        if (pf.space_before or 0) != 0:
            bad["before"] += 1
        if (pf.space_after or 0) != 0:
            bad["after"] += 1
        if pf.first_line_indent is None or round(pf.first_line_indent.mm, 1) != 12.5:
            bad["indent"] += 1
        if p.alignment != 3:
            bad["align"] += 1
    res["Основной текст"] = "PASS" if not any(bad.values()) else f"FAIL {bad} из {len(main)}"
    alltext = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    n_yo, n_dash = len(re.findall("[ёЁ]", alltext)), alltext.count("—")
    res["Символы ё, Ё, —"] = "PASS" if not (n_yo or n_dash) else f"FAIL ё/Ё {n_yo}, — {n_dash}"
    # латиница курсивом
    nonital = []
    def scan(p):
        if (p.style.name or "").lower().startswith(("heading", "toc")):
            return
        if "/" in p.text or "\\" in p.text or "@" in p.text:
            return
        for r in p.runs:
            if r.text and LATIN.search(r.text) and not r.italic:
                if r._r.find(qn("w:instrText")) is None:
                    nonital.append(r.text[:30])
    for p in doc.paragraphs:
        scan(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    scan(p)
    res["Латиница курсивом"] = "PASS" if not nonital else f"FAIL {len(nonital)}: {nonital[:3]}"
    # перечисления
    lists = [p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None]
    badl = [p for p in lists if (p.paragraph_format.first_line_indent is not None and p.paragraph_format.first_line_indent.mm > 0) or (p.paragraph_format.space_before or 0) or (p.paragraph_format.space_after or 0)]
    long_dash = [p for p in lists if p.text.strip().startswith("—")]
    res["Перечисления"] = ("PASS" if not badl and not long_dash else f"FAIL {len(badl)}") + f" (пунктов {len(lists)})"
    # таблицы
    tb = {"width": 0, "font": 0, "halign": 0, "valign": 0, "space": 0, "indent": 0, "header": 0}
    for t in doc.tables:
        tw = t._tbl.tblPr.find(qn("w:tblW"))
        wmm = int(tw.get(qn("w:w"))) / 56.7 if tw is not None and tw.get(qn("w:w")) else 0
        if abs(wmm - 165) > 1.5:
            tb["width"] += 1
        if t.rows[0]._tr.trPr is None or t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is None:
            tb["header"] += 1
        for row in t.rows:
            for c in row.cells:
                if c.vertical_alignment != 1:
                    tb["valign"] += 1
                for p in c.paragraphs:
                    if not p.text.strip():
                        continue
                    if p.alignment != 1:
                        tb["halign"] += 1
                    pf = p.paragraph_format
                    if (pf.space_before or 0) or (pf.space_after or 0):
                        tb["space"] += 1
                    if pf.first_line_indent is not None and pf.first_line_indent.mm > 0:
                        tb["indent"] += 1
                    if any(r.font.size != Pt(12) or (r.font.name or "Times New Roman") != "Times New Roman" for r in p.runs if r.text.strip()):
                        tb["font"] += 1
    res["Таблицы"] = "PASS" if not any(tb.values()) else f"FAIL {tb} (таблиц {len(doc.tables)})"
    # рисунки: подпись сразу после рисунка
    figs = 0
    badf = 0
    pars = doc.paragraphs
    for i, p in enumerate(pars):
        if p._p.findall(".//" + qn("w:drawing")):
            figs += 1
            nxt = pars[i + 1].text if i + 1 < len(pars) else ""
            if not re.match(r"^Рисунок [\d.]+ [–-]", nxt):
                badf += 1
    res["Рисунки"] = "PASS" if not badf else f"FAIL {badf} без подписи под рисунком"
    res["Рисунки"] += f" (рисунков {figs})"
    # заголовки
    heads = [(i, p) for i, p in enumerate(pars) if (p.style.name or "").lower().startswith("heading")]
    dots = [p.text for _, p in heads if p.text.strip().endswith(".")]
    twice = sum(1 for (i, _), (j, _) in zip(heads, heads[1:]) if j == i + 1)
    res["Заголовки"] = "PASS" if not dots else f"FAIL точка в конце: {dots[:2]}"
    res["Заголовки"] += f" (подряд без текста: {twice})"
    empt = sum(1 for k, p in enumerate(pars) if not p.text.strip() and not p._p.findall(".//" + qn("w:drawing")) and p._p.find(".//" + qn("w:br")) is None and p._p.find(".//" + qn("w:fldChar")) is None and k != len(pars) - 1)
    res["Пустые абзацы"] = "PASS" if empt == 0 else f"FAIL {empt}"
    dbl = len(re.findall(r"\S  +\S", "\n".join(p.text for p in pars)))
    res["Повторные пробелы"] = "PASS" if dbl == 0 else f"FAIL {dbl}"
    # ссылки на источники
    txt = "\n".join(p.text for p in pars)
    m = re.search(r"Список использованных источников(.*?)(?:\nПриложение|\Z)", txt, re.S | re.I)
    res["Источники"] = "MANUAL (двусторонняя сверка ссылок и содержания)"
    return res


if __name__ == "__main__":
    targets = sys.argv[1:] or sorted(glob.glob(os.path.join(ROOT, "ЛР*", "ОТЧЕТ_*.docx")) + glob.glob(os.path.join(ROOT, "Финальный_отчет", "ОТЧЕТ_*.docx")))
    for f in targets:
        print("==", os.path.relpath(f, ROOT))
        for k, v in check(f).items():
            print(f"  {k}: {v}")
