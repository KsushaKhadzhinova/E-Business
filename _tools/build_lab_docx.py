"""Сборка отчётов ЛР в Word строго по СТП БГУИР 01–2024.

Запуск:
    python _tools/build_lab_docx.py            # все ЛР + финальный отчёт
    python _tools/build_lab_docx.py ЛР4        # одна ЛР (по префиксу папки)

Для каждой папки ЛР берётся ОТЧЕТ.md и сохраняется ОТЧЕТ_<ЛР>_NotaCode.docx.
Оформление: A4, поля 30/15/20/20 мм, Times New Roman 14, межстрочный 18 пт,
абзацный отступ 1,25 см, выравнивание по ширине, заголовки разделов полужирные,
ненумерованные заголовки по центру, подписи «Рисунок N – …» под рисунком по центру,
«Таблица N – …» над таблицей слева, приложения с новой страницы,
номера страниц внизу справа, титульный лист без номера.
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent

STUDENT = "К. А. Хаджинова"
GROUP = "320604"
TEACHER = "🔲 И. О. Фамилия"
FACULTY = "информационных технологий и управления"
DEPARTMENT = "информационных технологий автоматизированных систем"
DISCIPLINE = "Электронный бизнес"
YEAR = "2026"

LABS = {
    "ЛР1_Спрос_и_границы_рынка": ("№1", "Анализ поискового спроса и фиксация границ рынка"),
    "ЛР2-3_Объем_рынка": ("№2–3", "Оценка объёма рынка электронного бизнеса"),
    "ЛР4_Пять_сил_Портера": ("№4", "Анализ пяти сил Портера и барьеров входа на рынок"),
    "ЛР5_Анализ_конкурентов": ("№5", "Анализ конкурентов по сайтам, моделям Левитта и Кано"),
    "ЛР6_ЦП_Канвас_сценарии_релизы": ("№6", "Ценностное предложение, бизнес-модель Канвас, сценарии, экраны и релизы"),
    "ЛР7_Итоговый_отчет": ("№7", "Итоговый отчёт по бизнес-анализу цифрового товара"),
}

UNNUMBERED = ("выводы", "заключение", "введение", "список использованных источников", "содержание")
CAPTION_FIG = re.compile(r"^Рисунок\s+[А-ЯA-Z]?\.?\d+(\.\d+)?\s*[–-]")
CAPTION_TAB = re.compile(r"^Таблица\s+[А-ЯA-Z]?\.?\d+(\.\d+)?\s*[–-]")
FORMULA = re.compile(r"^[^.]*=.*\(\d+\)\s*$")
APPENDIX = re.compile(r"^Приложение\s+([А-ЯA-Z])\s*\(([^)]+)\)\s*(.*)$", re.I)


def reference_docx(path: Path) -> None:
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout
    path.write_bytes(data)
    doc = Document(str(path))
    for st in doc.styles:
        try:
            font = st.font
        except AttributeError:
            continue
        font.name = "Times New Roman"
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        font.color.rgb = RGBColor(0, 0, 0)
    doc.save(str(path))


def md_to_docx(md_text: str, out: Path, resource: Path) -> None:
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        ref = td / "ref.docx"
        reference_docx(ref)
        src = td / "in.md"
        src.write_text(md_text, encoding="utf-8")
        subprocess.run(["pandoc", str(src), "-f", "markdown-yaml_metadata_block-implicit_figures+pipe_tables+raw_attribute",
                        "-o", str(out), "--reference-doc", str(ref), "--resource-path", str(resource)], check=True)


def prepare_md(text: str) -> str:
    text = re.sub(r"```mermaid.*?```", "*(схема приведена в виде рисунка)*", text, flags=re.S)
    lines = text.splitlines()
    start = next((i for i, l in enumerate(lines) if re.match(r"^##\s+1\s", l)), None)
    if start is None:
        start = next((i for i, l in enumerate(lines) if l.startswith("## ")), 0)
    body = "\n".join(lines[start:])
    body = re.sub(r"^(#{2,6}) ", lambda m: m.group(1)[1:] + " ", body, flags=re.M)
    body = re.sub(r"\[([^\]]+)\]\((?!http)[^)]*\.md[^)]*\)", r"\1", body)
    out, prev = [], ""
    for line in body.splitlines():
        is_list = bool(re.match(r"^\s*([-*+]|\d+[.)])\s", line))
        prev_list = bool(re.match(r"^\s*([-*+]|\d+[.)])\s", prev))
        if is_list and prev.strip() and not prev_list and not prev.startswith("|"):
            out.append("")
        if line.startswith("|") and prev.strip() and not prev.startswith("|"):
            out.append("")
        if prev.startswith("|") and line.strip() and not line.startswith("|"):
            out.append("")
        out.append(line)
        prev = line
    return "\n".join(out)


def set_par(p, *, align=None, indent=None, before=0, after=0, bold=None, size=14, caps=False, keep_next=False):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.first_line_indent = indent
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.keep_with_next = keep_next
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None:
            r.font.bold = bold
        if caps:
            r.text = r.text.upper()


def page_break_before(p):
    p.paragraph_format.page_break_before = True


def insert_before(anchor, text="", **fmt):
    new = anchor.insert_paragraph_before(text)
    set_par(new, **fmt)
    return new


def field_run(p, instr):
    run = p.add_run()
    for kind, val in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    return run


def title_page(doc, first, lab_no, topic, kind):
    rows = [
        ("Министерство образования Республики Беларусь", 0),
        ("", 0),
        ("Учреждение образования", 0),
        ("БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ", 0),
        ("ИНФОРМАТИКИ И РАДИОЭЛЕКТРОНИКИ", 0),
        ("", 0),
        (f"Факультет {FACULTY}", 0),
        ("", 0),
        (f"Кафедра {DEPARTMENT}", 0),
    ] + [("", 0)] * 8 + [
        ("ОТЧЕТ", 1),
        (f"по {kind} {lab_no}", 0),
        ("по дисциплине", 0),
        (f"«{DISCIPLINE}»", 0),
        (f"на тему «{topic}»", 0),
    ] + [("", 0)] * 5
    for text, bold in rows:
        insert_before(first, text, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), bold=bool(bold))
    for label, value in (("Выполнила:", f"ст. гр. {GROUP}"), ("", STUDENT), ("", ""), ("Проверил(а):", TEACHER)):
        p = insert_before(first, "", align=WD_ALIGN_PARAGRAPH.LEFT, indent=Mm(0))
        p.paragraph_format.left_indent = Mm(0)
        r1 = p.add_run(label)
        r1.font.size = Pt(14)
        p.add_run("\t" + value).font.size = Pt(14)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(115))
    for _ in range(7):
        insert_before(first, "", indent=Mm(0))
    insert_before(first, f"Минск {YEAR}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0))
    toc_head = insert_before(first, "СОДЕРЖАНИЕ", align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), bold=True, after=18)
    page_break_before(toc_head)
    toc = insert_before(first, "", indent=Mm(0))
    field_run(toc, 'TOC \\o "1-2" \\h \\z \\u')
    page_break_before(first)


def page_numbers(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(30), Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field_run(p, "PAGE")


def format_body(doc):
    body_pars = list(doc.paragraphs)
    for p in body_pars:
        style = (p.style.name or "").lower()
        text = p.text.strip()
        if style.startswith("heading"):
            level = int(re.sub(r"\D", "", style) or 1)
            m = APPENDIX.match(text)
            if m and level == 1:
                letter, kind, title = m.groups()
                for r in p.runs[1:]:
                    r.text = ""
                p.runs[0].text = f"ПРИЛОЖЕНИЕ {letter.upper()}"
                set_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), bold=True, size=14)
                page_break_before(p)
                k = p._p.addnext(OxmlElement("w:p"))
                nxt = p._p.getnext()
                from docx.text.paragraph import Paragraph
                kp = Paragraph(nxt, p._parent)
                kp.add_run(f"({kind.lower()})")
                set_par(kp, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0))
                if title:
                    kp._p.addnext(OxmlElement("w:p"))
                    tp = Paragraph(kp._p.getnext(), p._parent)
                    tp.add_run(title.strip())
                    set_par(tp, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), bold=True, after=18)
                continue
            if level == 1 and text.lower() in UNNUMBERED:
                set_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), bold=True, after=18, caps=True, keep_next=True)
                page_break_before(p)
            elif level == 1:
                set_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Mm(12.5), bold=True, after=18, caps=True, keep_next=True)
                page_break_before(p)
            else:
                set_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Mm(12.5), bold=True, before=18, after=18, keep_next=True)
            continue
        has_img = bool(p._p.findall(".//" + qn("w:drawing")))
        if has_img:
            set_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), before=12, keep_next=True)
        elif CAPTION_FIG.match(text):
            set_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), after=12)
        elif CAPTION_TAB.match(text):
            set_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Mm(0), before=12, keep_next=True)
        elif FORMULA.match(text) and len(text) < 140:
            set_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Mm(0), before=6, after=6)
        elif p._p.getprevious() is not None and p._p.getprevious().tag == qn("w:tbl"):
            set_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Mm(12.5), before=12)
        elif p._p.getparent().tag == qn("w:body"):
            set_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Mm(12.5))
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = t._tbl.tblPr
        for old in tbl_pr.findall(qn("w:tblBorders")) + tbl_pr.findall(qn("w:tblStyle")) + tbl_pr.findall(qn("w:tblW")) + tbl_pr.findall(qn("w:tblLayout")):
            tbl_pr.remove(old)
        fit_columns(t)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "000000")
            borders.append(el)
        tbl_pr.append(borders)
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    set_par(p, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT,
                            indent=Mm(0), size=12, bold=True if i == 0 else None)
                    p.paragraph_format.line_spacing = Pt(14)
    max_w = Mm(165)
    for shape in doc.inline_shapes:
        if shape.width > max_w:
            ratio = max_w / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)


def fit_columns(t, total_mm=165):
    ncols = len(t.columns)
    if ncols == 0 or not t.rows:
        return
    weights = [0.0] * ncols
    for row in t.rows:
        for j, cell in enumerate(row.cells[:ncols]):
            words = cell.text.split()
            longest = max((len(w) for w in words), default=1)
            weights[j] = max(weights[j], min(len(cell.text), 60) * 0.6 + longest)
    min_w = [12.0] * ncols
    for row in t.rows:
        for j, cell in enumerate(row.cells[:ncols]):
            lw = len(max((cell.text.split() or [""]), key=len))
            min_w[j] = max(min_w[j], min(lw * 2.4 + 3, 38.0))
    s = sum(weights) or 1
    widths = [max(total_mm * w / s, min_w[j]) for j, w in enumerate(weights)]
    k = total_mm / sum(widths)
    widths = [w * k for w in widths]
    tbl_pr = t._tbl.tblPr
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(total_mm * 56.7)))
    tw.set(qn("w:type"), "dxa")
    tbl_pr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tbl_pr.append(lay)
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if j < ncols:
            gc.set(qn("w:w"), str(int(widths[j] * 56.7)))
    for row in t.rows:
        for j, cell in enumerate(row.cells[:ncols]):
            cell.width = Mm(widths[j])


def build(md_text: str, resource: Path, out: Path, lab_no: str, topic: str, kind: str = "лабораторной работе"):
    md_to_docx(prepare_md(md_text), out, resource)
    doc = Document(str(out))
    page_numbers(doc)
    format_body(doc)
    title_page(doc, doc.paragraphs[0], lab_no, topic, kind)
    doc.save(str(out))
    print("OK", out.relative_to(ROOT))


def build_lab(folder: str):
    lab_no, topic = LABS[folder]
    src = ROOT / folder / "ОТЧЕТ.md"
    if not src.exists():
        print("нет ОТЧЕТ.md:", folder)
        return
    kind = "лабораторным работам" if "–" in lab_no else "лабораторной работе"
    short = folder.split("_")[0]
    build(src.read_text(encoding="utf-8"), ROOT / folder, ROOT / folder / f"ОТЧЕТ_{short}_NotaCode.docx", lab_no, topic, kind)


def build_final():
    parts = []
    intro = ROOT / "Финальный_отчет" / "00_Введение.md"
    if intro.exists():
        parts.append("## Введение\n\n" + prepare_md_intro(intro.read_text(encoding="utf-8")))
    n = 0
    for folder, (lab_no, topic) in LABS.items():
        src = ROOT / folder / "ОТЧЕТ.md"
        if not src.exists():
            continue
        n += 1
        body = src.read_text(encoding="utf-8")
        body = re.sub(r"\]\((?!http)(Рисунки/[^)]+)\)", lambda m: f"]({(ROOT / folder / m.group(1)).as_posix()})", body)
        body = "\n".join(body.splitlines()[next((i for i, l in enumerate(body.splitlines()) if l.startswith("## ")), 0):])
        body = re.sub(r"^(#{2,5}) ", lambda m: "#" + m.group(1) + " ", body, flags=re.M)
        body = re.sub(r"^(#{3,6}) (\d+(?:\.\d+)*) ", lambda m: f"{m.group(1)} {n}.{m.group(2)} ", body, flags=re.M)
        parts.append(f"## {n} Лабораторная работа {lab_no}. {topic}\n\n" + body)
    for extra in ("99_Заключение.md", "98_Список_источников.md"):
        f = ROOT / "Финальный_отчет" / extra
        if f.exists():
            parts.append(f.read_text(encoding="utf-8"))
    md = "\n\n".join(parts)
    (ROOT / "Финальный_отчет" / "Финальный_отчет.md").write_text(md, encoding="utf-8")
    build(md, ROOT, ROOT / "Финальный_отчет" / "ОТЧЕТ_по_лабораторным_работам_1-7_NotaCode.docx",
          "№1–7", "Бизнес-анализ цифрового продукта NotaCode", "лабораторным работам")


def prepare_md_intro(text):
    return re.sub(r"^# .*\n", "", text)


if __name__ == "__main__":
    args = sys.argv[1:]
    targets = [f for f in LABS if not args or any(f.startswith(a) for a in args)]
    for f in targets:
        build_lab(f)
    if not args or "final" in args:
        build_final()
