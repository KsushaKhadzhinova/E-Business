# -*- coding: utf-8 -*-
"""Углублённая проверка ЛР1-ЛР7: ссылки на таблицы и рисунки, согласованность ключевых чисел, шаблонные заглушки, содержимое docx и Excel."""
import glob
import os
import re
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LABS = ["ЛР1_Спрос_и_границы_рынка", "ЛР2-3_Объем_рынка", "ЛР4_Пять_сил_Портера", "ЛР5_Анализ_конкурентов", "ЛР6_ЦП_Канвас_сценарии_релизы", "ЛР7_Итоговый_отчет"]
T = {l: open(os.path.join(ROOT, l, "ОТЧЕТ.md"), encoding="utf-8").read().replace(" ", " ").replace(" ", " ") for l in LABS}
issues = []


def add(lab, msg):
    issues.append(f"{lab.split('_')[0]}: {msg}")


# 1. ссылки «таблица N / рисунок N» не выходят за диапазон
for lab, t in T.items():
    nt = len(re.findall(r"^Таблица \d+ – ", t, re.M))
    nf = len(re.findall(r"^Рисунок \d+ – ", t, re.M))
    body = re.sub(r"^Таблица \d+ – .*$", "", t, flags=re.M)
    for m in re.finditer(r"(?<![А-Яа-я])(?:табл\.|таблиц[аеуы]|таблицах)\s+(\d+)(?![\d.])", body, re.I):
        k = int(m.group(1))
        ctx = body[max(0, m.start() - 12):m.start()]
        if k > nt and "ЛР" not in ctx and "[" not in ctx:
            add(lab, f"ссылка на таблицу {k} при {nt} таблицах: …{body[max(0, m.start() - 30):m.end() + 5].strip()}")
    for m in re.finditer(r"(?<![А-Яа-я])рисун(?:ок|ке|ка|ку)\s+(\d+)(?![\d.])", body, re.I):
        k = int(m.group(1))
        if k > nf and "ЛР" not in body[max(0, m.start() - 8):m.start()]:
            add(lab, f"ссылка на рисунок {k} при {nf} рисунках")
# 2. заглушки
for lab, t in T.items():
    for pat in (r"\[[А-Яа-я ]{3,40}\]\s*$", r"\bTODO\b", r"Пример, замените", r"lorem", r"XXX", r"\?\?\?"):
        n = len(re.findall(pat, t, re.M | re.I))
        if n:
            add(lab, f"возможная заглушка «{pat}»: {n}")
    if re.search(r"[ \t]{3,}\S", t.replace("|", "")) and False:
        pass
# 3. ключевые числа
keys = {"40 USD": ["ЛР1", "ЛР2-3", "ЛР4", "ЛР5", "ЛР6", "ЛР7"], "49 893": ["ЛР2-3", "ЛР4", "ЛР7"], "414 754": ["ЛР2-3", "ЛР4", "ЛР7"], "16 590": ["ЛР2-3", "ЛР4", "ЛР6", "ЛР7"], "229,0": ["ЛР1", "ЛР4", "ЛР7"]}
for k, labs in keys.items():
    for lab in LABS:
        short = lab.split("_")[0]
        if short in labs and k not in T[lab]:
            add(lab, f"не найдено ключевое число «{k}»")
# 4. противоречия
for lab, t in T.items():
    if re.search(r"3,56\s*%", t) and lab.startswith("ЛР7"):
        add(lab, "устаревшая доля стипендии 3,56 %")
    if "Нотация недоступна в Free" in t and "Было в ЛР6" not in t:
        add(lab, "устаревшее утверждение о нотациях в Free")
# 5. docx
from docx import Document  # noqa: E402

for lab in LABS + ["Финальный_отчет"]:
    d = os.path.join(ROOT, lab)
    for f in glob.glob(os.path.join(d, "ОТЧЕТ*.docx")):
        doc = Document(f)
        name = os.path.basename(f)
        nt = len(doc.tables)
        ni = len(doc.inline_shapes)
        if lab != "Финальный_отчет":
            mdt = len(re.findall(r"^\|(?:\s*:?-{3,}:?\s*\|)+\s*$", T[lab], re.M))
            if abs(nt - mdt) > 1:
                add(lab, f"{name}: таблиц в docx {nt}, в md {mdt}")
            mdi = len(re.findall(r"!\[\]\(", T[lab]))
            if ni != mdi:
                add(lab, f"{name}: рисунков в docx {ni}, в md {mdi}")
        sec = doc.sections[0]
        mm = lambda v: round(v.mm)
        if (mm(sec.left_margin), mm(sec.right_margin), mm(sec.top_margin), mm(sec.bottom_margin)) != (30, 15, 20, 20):
            add(lab, f"{name}: поля {mm(sec.left_margin)}/{mm(sec.right_margin)}/{mm(sec.top_margin)}/{mm(sec.bottom_margin)} мм вместо 30/15/20/20")
        normal = doc.styles["Normal"].font
        body_fonts = {r.font.name for p in doc.paragraphs[:60] for r in p.runs if r.font.name}
        if normal.name != "Times New Roman" and "Times New Roman" not in body_fonts:
            add(lab, f"{name}: шрифт {normal.name}")
        empty = sum(1 for p in doc.paragraphs if re.search(r"🔲\s*$", p.text))
# 6. Excel: заглушки шаблонов и «Пример»
for x in glob.glob(os.path.join(ROOT, "ЛР[56]*", "Excel", "*.xlsx")):
    z = zipfile.ZipFile(x)
    ss = z.read("xl/sharedStrings.xml").decode("utf-8", "ignore") if "xl/sharedStrings.xml" in z.namelist() else ""
    n = ss.count("Пример, замените") + ss.count("Пример.")
    if n:
        add(os.path.basename(os.path.dirname(os.path.dirname(x))), f"{os.path.basename(x)}: остались образцы шаблона ({n})")
# 7. ссылки на файлы в приложениях
for lab, t in T.items():
    for m in re.finditer(r"`?((?:Excel|Python|Рисунки|Материалы_собранные)/[A-Za-z0-9_./-]+\.(?:xlsx|py|png|json|csv))`?", t):
        pth = os.path.join(ROOT, lab, m.group(1))
        if not os.path.exists(pth):
            add(lab, f"в отчёте упомянут отсутствующий файл {m.group(1)}")
print("\n".join(issues) if issues else "замечаний нет")
print("всего замечаний:", len(issues))
