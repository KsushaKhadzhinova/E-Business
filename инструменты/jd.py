import openpyxl,os
n=0
for root,_,fs in os.walk("СДАЧА"):
    for f in fs:
        if f.endswith(".xlsx") and "Журнал_выгрузки" in openpyxl.load_workbook(os.path.join(root,f),read_only=True).sheetnames:
            ws=openpyxl.load_workbook(os.path.join(root,f))["Журнал_выгрузки"]
            vals=[(c.value,type(c.value).__name__) for c in ws["A"][1:6] if c.value is not None]
            bad=[v for v in vals if v[1]!="str" ]
            print(f[:34],vals[:2],"BAD" if bad else "ok")
