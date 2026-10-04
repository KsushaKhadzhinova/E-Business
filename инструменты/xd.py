import openpyxl,warnings,collections,os
warnings.filterwarnings("ignore")
T=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad"
f="XLT_YW_dfd_diagramma_BY.xlsx"
a=openpyxl.load_workbook(os.path.join(T,"broken",f)); b=openpyxl.load_workbook(os.path.join(r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР1",f))
cnt=collections.Counter(); ex=collections.defaultdict(list)
for ws in a.worksheets:
    w2=b[ws.title]
    for row in ws.iter_rows():
        for c in row:
            o=w2[c.coordinate]
            if c.value!=o.value:
                k=(ws.title,"val"); cnt[k]+=1
                if len(ex[k])<3: ex[k].append((c.coordinate,repr(c.value)[:60],repr(o.value)[:60]))
            elif c.number_format!=o.number_format:
                k=(ws.title,"fmt"); cnt[k]+=1
                if len(ex[k])<2: ex[k].append((c.coordinate,c.number_format,o.number_format))
for k,v in cnt.items(): print(k,v,ex[k])
