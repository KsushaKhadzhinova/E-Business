import openpyxl,os,warnings,collections,json,datetime
warnings.filterwarnings("ignore")
T=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad"
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
def tx(v): return getattr(v,"text",v)
tot=collections.Counter(); other=[]
for root,_,fs in os.walk(base):
    for f in sorted(fs):
        if not f.endswith(".xlsx"): continue
        a=openpyxl.load_workbook(os.path.join(T,"broken",f)); b=openpyxl.load_workbook(os.path.join(root,f)); ch=[]
        for ws in a.worksheets:
            w2=b[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    o=w2[c.coordinate]; av,bv=tx(c.value),tx(o.value)
                    if av!=bv:
                        if isinstance(av,str) and av.startswith("="):
                            if isinstance(bv,str) and av.replace("\\","")==bv: k="formula_backslash"
                            else:
                                k="formula_other"; other.append((f[:25],ws.title,c.coordinate,str(av)[:70],str(bv)[:70]))
                        elif isinstance(av,str) and isinstance(bv,str) and av.rstrip(".")==bv.rstrip("."): k="text_trailing_dot"
                        else: k="text_other"; other.append((f[:25],ws.title,c.coordinate,str(av)[:70],str(bv)[:70]))
                        tot[k]+=1
                    elif c.number_format!=o.number_format: tot["fmt"]+=1
print(dict(tot))
import random; random.seed(1)
for x in [o for o in other][:0]: pass
fo=[o for o in other]
print(len(fo)); 
for o in fo[:25]: print(o)
json.dump(fo,open(T+r"\other.json","w",encoding="utf-8"),ensure_ascii=False)
