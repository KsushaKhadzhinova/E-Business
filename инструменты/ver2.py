import openpyxl,os,warnings
warnings.filterwarnings("ignore")
p=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР5\klassifikaciya_urovni_konkurencii_NotaCode.xlsx"
wb=openpyxl.load_workbook(p,data_only=True)
for sh,rng in (("07_Карта",4),("04_Оценка",4),("06_Сводка",19)):
    ws=wb[sh]; print(sh,[ws.cell(rng,c).value for c in range(1,12)])
box=chr(0x1F532); err=0; bx=0; xl=0
for r,_,fs in os.walk("СДАЧА"):
    for f in fs:
        if f.endswith(".xlsx") and not f.startswith("~"):
            a=openpyxl.load_workbook(os.path.join(r,f)); b=openpyxl.load_workbook(os.path.join(r,f),data_only=True)
            for ws in a.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        v=getattr(c.value,"text",c.value)
                        if isinstance(v,str):
                            if box in v: bx+=1
                            if "XLOOKUP" in v: xl+=1
                        d=b[ws.title][c.coordinate].value
                        if isinstance(d,str) and d.startswith("#"): err+=1
print("boxes left",bx,"xlookup left",xl,"error cells",err)
