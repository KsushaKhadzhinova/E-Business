import openpyxl,os,re,warnings
from docx import Document
warnings.filterwarnings("ignore")
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
def P(f): return [os.path.join(r,f) for r,_,fs in os.walk(base) if f in fs][0]
def text(lr):
    d=Document(os.path.join(base,lr,f"ОТЧЕТ_{lr}_NotaCode.docx")); t=[p.text for p in d.paragraphs]
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells: t.append(c.text)
    return re.sub(r"[\s\xa0]+"," "," ".join(t))
def var(v):
    out=set()
    for dec in (0,1,2):
        s=f"{v:,.{dec}f}".replace(",", " ").replace(".",","); out.add(s); out.add(s.replace(" ",""))
        if abs(v)<=1.0001: out.add(f"{v*100:.{dec}f}".replace(".",",")+" %"); out.add(f"{v*100:.{dec}f}".replace(".",",")+"%")
    return out
sets={"ЛР5":[("klassifikaciya_urovni_konkurencii_NotaCode.xlsx","06_Сводка"),("analiz_konkurentov_Levitt_Kano_NotaCode.xlsx","09_Дашборд"),("analiz_biznes_modeli_konkurentov_NotaCode.xlsx","12_Сводка")],
      "ЛР6":[("6_1_cennostnoe_predlozhenie_NotaCode.xlsx","Дашборд"),("6_2_biznes_model_kanvas_NotaCode.xlsx","11_Дашборд"),("6_3_biznes_logika_scenarii_NotaCode.xlsx","Дашборд"),("6_4_stranicy_ekrany_NotaCode.xlsx","Дашборд"),("6_5_relizy_NotaCode.xlsx","Дашборд")]}
for lr,lst in sets.items():
    tx=text(lr)
    for f,sh in lst:
        try: ws=openpyxl.load_workbook(P(f),data_only=True)[sh]
        except KeyError: print(lr,f[:22],"no sheet",sh); continue
        vals=[(c.coordinate,c.value) for row in ws.iter_rows() for c in row if isinstance(c.value,(int,float)) and not isinstance(c.value,bool) and (abs(c.value)>=10 or (c.value!=int(c.value)))]
        hit=0; miss=[]
        for c,v in vals[:40]:
            if any(x in tx for x in var(v)): hit+=1
            else: miss.append((c,round(v,3)))
        print(lr,f[:22],sh,"found",hit,"of",min(len(vals),40),"missing",miss[:10])
