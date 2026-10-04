import openpyxl,os,json,warnings,re
warnings.filterwarnings("ignore")
T=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad"
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
def P(f): return [os.path.join(r,f) for r,_,fs in os.walk(base) if f in fs][0]
items=[]
def add(p,s,c,v): items.append({"p":p,"s":s,"c":c,"v":v})
# a/b: 6_4
p=P("6_4_stranicy_ekrany_NotaCode.xlsx"); wb=openpyxl.load_workbook(p,data_only=True)
names={r[0].value:r[2].value for r in wb["Страницы_и_экраны"].iter_rows(min_row=5) if r[0].value}
rows=[[c.value for c in r] for r in wb["Матрица_сценарий_интерфейс"].iter_rows(min_row=5,max_row=wb["Матрица_сценарий_интерфейс"].max_row) if r[0].value]
by={}
for i,r in enumerate(rows): by.setdefault(r[1],[]).append((int(r[5]),r[2],i))
for sc,l in by.items():
    l.sort()
    for k,(o,u,i) in enumerate(l):
        row=5+i; cur=rows[i]
        if cur[6] in (None,"") and k>0: add(p,"Матрица_сценарий_интерфейс",f"G{row}",l[k-1][1])
        if cur[7] in (None,"") and k<len(l)-1: add(p,"Матрица_сценарий_интерфейс",f"H{row}",l[k+1][1])
        if cur[8] in (None,""):
            nm=names.get(u,u); role="основной шаг" if cur[3]=="Основной" else "вспомогательный шаг"
            add(p,"Матрица_сценарий_интерфейс",f"I{row}",f"Шаг {o} сценария {sc}: {nm}; {role}, {str(cur[4]).lower()}")
ws=wb["Страницы_и_экраны"]; mp={"Принято":"Оставить","Требует уточнения":"Уточнить","Гипотеза":"Перенести в следующий выпуск"}
for r in range(5,ws.max_row+1):
    if ws.cell(r,1).value and not ws.cell(r,21).value and ws.cell(r,18).value in mp: add(p,"Страницы_и_экраны",f"U{r}",mp[ws.cell(r,18).value])
# c: Levitt 07/08
p=P("analiz_konkurentov_Levitt_Kano_NotaCode.xlsx"); wb=openpyxl.load_workbook(p,data_only=True); ws=wb["07_Стандарт_рынка"]
chk={"Рыночный стандарт":"Подтвердить реализацию в продуктах конкурентов, а не только на сайтах","Формирующаяся норма":"Отследить динамику распространенности признака","Низкий приоритет":"Пересмотреть при росте распространенности признака","Переходит в ожидаемое":"Оценить готовность пользователей платить за признак","Зона дифференциации":"Проверить спрос и стоимость реализации признака"}
for r in range(7,ws.max_row+1):
    if not ws.cell(r,1).value: continue
    st=ws.cell(r,9).value; n=ws.cell(r,6).value; sl=ws.cell(r,8).value
    if not ws.cell(r,11).value: add(p,"07_Стандарт_рынка",f"K{r}",chk.get(st,"Уточнить по данным конкурентов"))
    if not ws.cell(r,12).value:
        add(p,"07_Стандарт_рынка",f"L{r}",(f"Признак есть у {n} из 8 конкурентов (просмотр 01.10.2026)" if n else "Не найден на просмотренных страницах 8 конкурентов (01.10.2026)"))
ws=wb["08_Преимущества"]; use={"Сильное преимущество":"Использовать как ориентир при формулировании ценностного предложения","Потенциальное преимущество":"Учесть при развитии продукта после выпуска минимальной версии"}
for r in range(7,ws.max_row+1):
    if ws.cell(r,1).value and not ws.cell(r,14).value and ws.cell(r,13).value in use: add(p,"08_Преимущества",f"N{r}",use[ws.cell(r,13).value])
# d: port.json (old updated texts) for existing empty cells
for it in json.load(open(T+r"\port.json",encoding="utf-8")).values():
    for x in it: items.append({"p":x["p"],"s":x["s"],"c":x["c"],"v":x["v"]})
# de-dup (derived first)
seen=set(); out=[]
for x in items:
    k=(x["p"],x["s"],x["c"])
    if k in seen: continue
    seen.add(k); out.append(x)
json.dump(out,open(T+r"\fill3.json","w",encoding="utf-8"),ensure_ascii=False)
print(len(out))
