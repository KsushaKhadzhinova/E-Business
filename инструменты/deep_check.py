import openpyxl
import os

base = r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"

issues = []

for root, dirs, files in os.walk(base):
    for f in sorted(files):
        if not f.endswith('.xlsx'):
            continue
        path = os.path.join(root, f)
        rel = os.path.relpath(path, base)
        try:
            wb = openpyxl.load_workbook(path, data_only=False)
            for sname in wb.sheetnames:
                ws = wb[sname]
                for row in ws.iter_rows():
                    for cell in row:
                        v = cell.value
                        if v is None:
                            continue
                        # Формулы не должны стать строками
                        if isinstance(v, str) and v.startswith('='):
                            pass  # формула в виде строки - ок
                        # Проверяем: в Журнал_выгрузки col A должны быть строки DD.MM.YYYY
                        if sname == 'Журнал_выгрузки' and cell.column == 1 and cell.row > 1:
                            if isinstance(v, int) and v > 40000:
                                issues.append(f"{rel}|{sname}|{cell.coordinate}|serial={v}")
                            elif hasattr(v, 'strftime') and cell.number_format not in ('@', 'General'):
                                issues.append(f"{rel}|{sname}|{cell.coordinate}|datetime,fmt={cell.number_format}")
                        # Проверяем Ввод_данных col A - должны быть datetime с форматом yyyy\-mm\-dd
                        if sname == 'Ввод_данных' and cell.column == 1 and cell.row > 1:
                            if isinstance(v, str) and len(v) == 10 and '.' in v:
                                issues.append(f"{rel}|{sname}|{cell.coordinate}|СТАЛ СТРОКОЙ: {v!r} (формулы Расчет сломаны)")
            wb.close()
        except Exception as e:
            issues.append(f"CORRUPT: {rel} -> {e}")

if issues:
    print(f"НАЙДЕНО ПРОБЛЕМ: {len(issues)}")
    for i in issues[:50]:
        print(" ", i)
else:
    print("Проблем не найдено — типы ячеек в норме.")
