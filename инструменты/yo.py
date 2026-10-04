import openpyxl,os,warnings,collections
from docx import Document
warnings.filterwarnings("ignore")
yo=("\u0451","\u0401")
tot=collections.Counter(); ex={}
for r,_,fs in os.walk("СДАЧА"):
    for f in fs:
        if f.startswith("~"): continue
        p=os.path.join(r,f)
        if f.endswith(".xlsx"):
            wb=openpyxl.load_workbook(p)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        v=getattr(c.value,"text",c.value)
                        if isinstance(v,str) and any(y in v for y in yo): tot[f]+=1; ex.setdefault(f,(ws.title,c.coordinate,v[:60]))
        elif f.endswith(".docx"):
            d=Document(p); n=sum(1 for x in d.paragraphs if any(y in x.text for y in yo))
            for tb in d.tables:
                for rw in tb.rows:
                    for c in rw.cells:
                        if any(y in c.text for y in yo): n+=1
            if n: tot[f]=n
print(dict(tot)); print(ex)
for f in ["klassifikaciya_urovni_konkurencii_NotaCode.xlsx","analiz_konkurentov_Levitt_Kano_NotaCode.xlsx"]:
    for r,_,fs in os.walk("СДАЧА"):
        if f in fs:
            wb=openpyxl.load_workbook(os.path.join(r,f)); print(f[:30],{ws.title:ws.max_row for ws in wb.worksheets})
