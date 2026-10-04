from docx import Document
from docx.oxml.ns import qn
import glob,os,openpyxl,warnings
warnings.filterwarnings("ignore")
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
for p in glob.glob(base+r"\**\*.docx",recursive=True):
    d=Document(p); n=0
    for t in d.element.iter(qn("w:t")):
        if t.text and "320604" in t.text: t.text=t.text.replace("320604","60131"); n+=1
    for s in d.sections:
        for part in (s.header,s.footer):
            try:
                for t in part._element.iter(qn("w:t")):
                    if t.text and "320604" in t.text: t.text=t.text.replace("320604","60131"); n+=1
            except Exception: pass
    if n: d.save(p)
    print(os.path.relpath(p,base),n)
for p in glob.glob(base+r"\**\*.xlsx",recursive=True):
    if os.path.basename(p).startswith("~"): continue
    wb=openpyxl.load_workbook(p,read_only=True); c=0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v,str) and "320604" in v: c+=1
    if c: print("xlsx",os.path.basename(p),c)
print("scan done")
