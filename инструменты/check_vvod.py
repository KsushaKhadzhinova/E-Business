import openpyxl
import os
from datetime import datetime, date

base = r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"

print("=== Ввод_данных col A — типы ячеек ===\n")

for root, dirs, files in os.walk(base):
    for f in sorted(files):
        if not f.endswith('.xlsx'):
            continue
        path = os.path.join(root, f)
        rel = os.path.relpath(path, base)
        try:
            wb = openpyxl.load_workbook(path, data_only=False)
            if 'Ввод_данных' not in wb.sheetnames:
                wb.close()
                continue
            ws = wb['Ввод_данных']
            bad = []
            for row in ws.iter_rows(min_col=1, max_col=1):
                cell = row[0]
                if cell.row == 1 or cell.value is None:
                    continue
                v = cell.value
                t = type(v).__name__
                fmt = cell.number_format
                if isinstance(v, str):
                    bad.append(f"  {cell.coordinate}: СТРОКА={v!r}  ← СЛОМАНО (формулы не считаются)")
                elif isinstance(v, (datetime, date)):
                    pass  # OK
                elif isinstance(v, int) and v > 40000:
                    bad.append(f"  {cell.coordinate}: serial_int={v}  fmt={fmt}")
            if bad:
                print(f"ПРОБЛЕМЫ в {rel}:")
                for b in bad:
                    print(b)
                print()
            else:
                print(f"OK  {rel}")
            wb.close()
        except Exception as e:
            print(f"ERROR {rel}: {e}")

print("\n=== Журнал_выгрузки col A — типы ячеек ===\n")

for root, dirs, files in os.walk(base):
    for f in sorted(files):
        if not f.endswith('.xlsx'):
            continue
        path = os.path.join(root, f)
        rel = os.path.relpath(path, base)
        try:
            wb = openpyxl.load_workbook(path, data_only=False)
            if 'Журнал_выгрузки' not in wb.sheetnames:
                wb.close()
                continue
            ws = wb['Журнал_выгрузки']
            sample = []
            for row in ws.iter_rows(min_col=1, max_col=1):
                cell = row[0]
                if cell.row == 1 or cell.value is None:
                    continue
                v = cell.value
                t = type(v).__name__
                fmt = cell.number_format
                sample.append(f"  {cell.coordinate}: {t}={v!r}  fmt={fmt}")
                if len(sample) >= 3:
                    break
            print(f"{rel}:")
            for s in sample:
                print(s)
            wb.close()
        except Exception as e:
            print(f"ERROR {rel}: {e}")
