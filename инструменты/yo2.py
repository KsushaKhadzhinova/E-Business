import openpyxl,os,json,warnings
from docx import Document
from docx.oxml.ns import qn
warnings.filterwarnings("ignore")
M={"\u0451":"\u0435","\u0401":"\u0415","\u2014":"\u2013"}
def fx(s):
    for k,v in M.items(): s=s.replace(k,v)
    return s
items=[]; base=os.path.abspath("СДАЧА")
for r,_,fs in os.walk(base):
    for f in fs:
        if f.startswith("~"): continue
        p=os.path.join(r,f)
        if f.endswith(".xlsx"):
            wb=openpyxl.load_workbook(p)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        v=getattr(c.value,"text",c.value)
                        if isinstance(v,str) and any(k in v for k in M):
                            items.append({"p":p,"s":ws.title,"c":c.coordinate,"v":fx(v),"f":v.startswith("=")})
        elif f.endswith(".docx"):
            d=Document(p); n=0
            for t in d.element.iter(qn("w:t")):
                if t.text and any(k in t.text for k in M): t.text=fx(t.text); n+=1
            for s in d.sections:
                for part in (s.header,s.footer):
                    try:
                        for t in part._element.iter(qn("w:t")):
                            if t.text and any(k in t.text for k in M): t.text=fx(t.text); n+=1
                    except Exception: pass
            if n: d.save(p); print("docx",f,n)
json.dump(items,open(r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad\yo.json","w",encoding="utf-8"),ensure_ascii=False)
print("xlsx cells to fix:",len(items))
