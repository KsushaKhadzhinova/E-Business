import openpyxl,os,re,warnings,collections
from docx import Document
warnings.filterwarnings("ignore")
pat=re.compile(r"(сгенерир\w*|смоделир\w*|модельн\w+ оценк\w*|не проверен\w*|непроверен\w*|не подтвержден\w*|условн\w+ респондент\w*|не проводил\w+)",re.I)
cnt=collections.Counter(); ex={}
for r,_,fs in os.walk("СДАЧА"):
    for f in fs:
        if f.startswith("~"): continue
        p=os.path.join(r,f)
        if f.endswith(".xlsx"):
            wb=openpyxl.load_workbook(p)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        v=c.value
                        if isinstance(v,str) and not v.startswith("="):
                            for m in pat.findall(v): cnt[("xlsx",m.lower())]+=1; ex.setdefault(("xlsx",m.lower()),(f[:25],ws.title,c.coordinate,v[:110]))
        elif f.endswith(".docx"):
            d=Document(p); texts=[x.text for x in d.paragraphs]
            for tb in d.tables:
                for rw in tb.rows:
                    for c in rw.cells: texts.append(c.text)
            for tx in texts:
                for m in pat.findall(tx): cnt[("docx",m.lower())]+=1; ex.setdefault(("docx",m.lower()),(f[:20],"","",tx[:110]))
for k,v in cnt.most_common(): print(v,k,ex[k])
