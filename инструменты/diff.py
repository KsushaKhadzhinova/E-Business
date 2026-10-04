import openpyxl,os,warnings,collections
warnings.filterwarnings("ignore")
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"; bro=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad\broken"
for root,_,fs in os.walk(base):
    for f in sorted(fs):
        if not f.endswith(".xlsx") or f.startswith("XLT"): continue
        a=openpyxl.load_workbook(os.path.join(bro,f)); b=openpyxl.load_workbook(os.path.join(root,f))
        diffs=collections.Counter(); ex=[]
        for ws in a.worksheets:
            if ws.title not in b.sheetnames: diffs["missing sheet "+ws.title]+=1; continue
            wb_=b[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    o=wb_[c.coordinate]
                    if c.value!=o.value:
                        diffs[ws.title+":value"]+=1
                        if len(ex)<3: ex.append((ws.title,c.coordinate,repr(c.value)[:30],repr(o.value)[:30]))
                    elif c.number_format!=o.number_format:
                        diffs[ws.title+":fmt"]+=1
                        if len(ex)<3: ex.append((ws.title,c.coordinate,c.number_format,o.number_format))
        print(f[:40],dict(diffs),ex[:3])
