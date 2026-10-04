import openpyxl,warnings,os
warnings.filterwarnings("ignore")
T=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad"
f="klassifikaciya_urovni_konkurencii_NotaCode.xlsx"
for label,p in (("CURRENT",r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР5\\"+f),("OLD-broken",T+"\\broken\\"+f)):
    wb=openpyxl.load_workbook(p); print("=====",label,wb.sheetnames)
    ws=wb["01_Компании"]
    for r in (3,4,5): print([str(c.value)[:22] if c.value is not None else None for c in ws[r][:8]])
    ws=wb["07_Карта"]
    for r in (3,4): print("07",[str(c.value)[:40] if c.value is not None else None for c in ws[r][:11]])
    ws=wb["04_Оценка"]
    for r in (3,4): print("04",[str(c.value)[:30] if c.value is not None else None for c in ws[r][:5]])
