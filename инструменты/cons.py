from docx import Document
import os,re
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
keys=["197 988","49 893","414 754","1 634","3,0285","121,14","2,32","3,12","60131","320604","ДОСНЯТЬ","модельная оценка","июнь 2026","август 2026","Similarweb"]
def alltext(d):
    t=[p.text for p in d.paragraphs]
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells: t.append(c.text)
    return "\n".join(t)
print("%-8s"%"", *["%-9s"%k[:9] for k in keys])
for lr in ["ЛР1","ЛР2-3","ЛР4","ЛР5","ЛР6","ЛР7"]:
    d=Document(os.path.join(base,lr,f"ОТЧЕТ_{lr}_NotaCode.docx")); t=alltext(d)
    print("%-8s"%lr, *["%-9d"%t.count(k) for k in keys])
