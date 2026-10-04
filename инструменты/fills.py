import openpyxl,os,warnings,collections,json
warnings.filterwarnings("ignore")
T=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad"
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
fills=collections.defaultdict(list); tot=collections.Counter(); vals=collections.Counter()
for root,_,fs in os.walk(base):
    for f in sorted(fs):
        if not f.endswith(".xlsx") or f.startswith("~"): continue
        a=openpyxl.load_workbook(os.path.join(T,"broken",f)); b=openpyxl.load_workbook(os.path.join(root,f))
        for ws in a.worksheets:
            if ws.title not in b.sheetnames: continue
            w2=b[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    v=c.value; o=w2[c.coordinate].value
                    if v is None or (isinstance(v,str) and (v.startswith("=") or "\\$" in v)): continue
                    if o is None or (isinstance(o,str) and o.strip()==""):
                        fills[f].append((ws.title,c.coordinate,v if not hasattr(v,"isoformat") else v.isoformat())); tot[f]+=1; vals[str(v)[:30]]+=1
print(dict(tot)); print(vals.most_common(15))
json.dump({k:[list(x) for x in v] for k,v in fills.items()},open(T+r"\fills.json","w",encoding="utf-8"),ensure_ascii=False,default=str)
