from docx import Document
import re
d=Document(r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР7\ОТЧЕТ_ЛР7_NotaCode.docx")
txt=[p.text for p in d.paragraphs]
for tb in d.tables:
    for r in tb.rows:
        for c in r.cells: txt.append(c.text)
seen=set()
for t in txt:
    for k in ["2,32","3,12","модельн","проверен","3,7","60131"]:
        if k in t and t not in seen:
            seen.add(t); i=t.find(k); print("["+k+"]",t[max(0,i-110):i+110].replace("\n"," "))
