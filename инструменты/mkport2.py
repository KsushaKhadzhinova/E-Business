import openpyxl,os,warnings,re,json,collections
warnings.filterwarnings("ignore")
T=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad"
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
names=["analiz_biznes_modeli_konkurentov_NotaCode.xlsx","analiz_konkurentov_Levitt_Kano_NotaCode.xlsx","klassifikaciya_urovni_konkurencii_NotaCode.xlsx","6_2_biznes_model_kanvas_NotaCode.xlsx","6_3_biznes_logika_scenarii_NotaCode.xlsx"]
ABBR={"у.е","т.д","т.п","т.е","мес","тыс","млн","руб","шт","стр","рис","табл","др"}
def clean(v):
    if not isinstance(v,str): return v
    v=v.replace("\u2014","\u2013").replace("\u0451","\u0435").replace("\u0401","\u0415").replace("Хаджынова К.А","К. А. Хаджинова")
    v=v.replace(chr(0x1F532)+" ДОСНЯТЬ","не собрано").replace(chr(0x1F532)+" ","")
    if v.endswith(".") and not v.endswith("..") and len(v)>3:
        m=re.search(r"([A-Za-zА-Яа-я]+)$",v[:-1]); last=m.group(1).lower() if m else ""
        if not (last in ABBR and len(last)<=4) and not v[:-1].endswith("у.е"): v=v[:-1]
    return v
isid=lambda x: isinstance(x,(int,float)) or (isinstance(x,str) and re.match(r"^[A-Za-zА-Я]{0,4}[-_]?\d{1,4}$",x.strip()) is not None)
out=[]
for root,_,fs in os.walk(base):
    for f in fs:
        if f in names:
            p=os.path.join(root,f)
            a=openpyxl.load_workbook(os.path.join(T,"broken",f)); b=openpyxl.load_workbook(os.path.join(T,"prefinal",f))
            n=0
            for ws in a.worksheets:
                if ws.title not in b.sheetnames: continue
                w2=b[ws.title]; rows=collections.defaultdict(list)
                for row in ws.iter_rows():
                    for c in row:
                        v=c.value; o=w2[c.coordinate].value
                        if v is None or hasattr(v,"isoformat") or (isinstance(v,str) and (v.startswith("=") or "\\$" in v or v.strip() in ("","□"))) or hasattr(v,"text"): continue
                        if o is None or (isinstance(o,str) and o.strip()==""): rows[c.row].append((c.coordinate,v))
                for r,l in rows.items():
                    if all(isid(v) for _,v in l): continue
                    for coord,v in l: out.append({"p":p,"s":ws.title,"c":coord,"v":clean(v)}); n+=1
            print(f[:42],n)
json.dump(out,open(T+r"\port2.json","w",encoding="utf-8"),ensure_ascii=False)
print("total",len(out))
