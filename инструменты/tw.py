from docx import Document
import re
d=Document(r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР2-3\ОТЧЕТ_ЛР2-3_NotaCode.docx")
sec=d.sections[0]; print("text width cm",round((sec.page_width-sec.left_margin-sec.right_margin)/360000,1))
for i in (0,5,20,40):
    x=d.tables[i]._tbl.xml
    print(i,re.findall(r'<w:tblW[^>]*/>',x),re.findall(r'<w:tblLayout[^>]*/>',x),re.findall(r'<w:gridCol[^>]*/>',x)[:6],re.findall(r'<w:tblCellMar>.*?</w:tblCellMar>',x,re.S)[:1],len(d.tables[i].columns))
