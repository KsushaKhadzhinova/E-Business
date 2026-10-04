from docx import Document
from docx.text.paragraph import Paragraph

p = r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР4\ОТЧЕТ_ЛР4_NotaCode.docx"
d = Document(p)
rules = {
    "Уровень по шаблону": "Низкие барьеры (2,32 меньше порога 2,5)",
    "Справочно: индекс, поделенный на сумму весов": "2,23",
    "Уровень при справочном расчете": "Низкие барьеры",
}
n = 0
for tb in d.tables:
    first = tb.rows[0].cells[0].text.strip()
    if not first.startswith("Показатель"):
        continue
    labels = [r.cells[0].text for r in tb.rows]
    if not any("Индекс барьеров по шаблону" in x for x in labels):
        continue
    for r in tb.rows:
        key = r.cells[0].text
        for k, v in rules.items():
            if key.startswith(k):
                cell = r.cells[1]
                par = cell.paragraphs[0]
                if par.runs:
                    par.runs[0].text = v
                    for extra in par.runs[1:]:
                        extra._r.getparent().remove(extra._r)
                else:
                    par.add_run(v)
                n += 1
                print("fixed:", key[:60], "->", v)
d.save(p)
print("cells fixed", n)
