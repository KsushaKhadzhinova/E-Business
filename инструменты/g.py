from docx import Document
d=Document(r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР5\ОТЧЕТ_ЛР5_NotaCode.docx")
ps=d.paragraphs
for i,p in enumerate(ps):
    if p.text.startswith("1 ЦЕЛЬ"):
        for q in ps[i-2:i+4]: print(q.style.name,"|",q.text[:300],"| pb:",q.paragraph_format.page_break_before)
        break
d=Document(r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР1\ОТЧЕТ_ЛР1_NotaCode.docx")
ps=d.paragraphs
for i,p in enumerate(ps):
    if p.text.startswith("1 ХОД"):
        for q in ps[i-3:i+3]: print(q.style.name,"|",q.text[:100],"| pb:",q.paragraph_format.page_break_before)
        break
print([r.text for r in p.runs])
