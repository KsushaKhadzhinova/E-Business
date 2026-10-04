from docx import Document
from collections import Counter
d=Document(r"E:\ИИТ\ЭлектронныйБизнес\ПОЛНЫЕ_ВЕРСИИ_до_сокращения\ОТЧЕТ_ЛР2-3_полный.docx")
print("tables",len(d.tables),"paras",len(d.paragraphs),"empty",sum(1 for p in d.paragraphs if not p.text.strip()),"inline imgs",len(d.inline_shapes))
sz=Counter();ls=Counter();sa=Counter()
for tb in d.tables:
    for r in tb.rows:
        for c in r.cells:
            for p in c.paragraphs:
                pf=p.paragraph_format
                ls[str(pf.line_spacing)]+=1; sa[str(pf.space_after)]+=1
                for run in p.runs: sz[str(run.font.size.pt if run.font.size else None)]+=1
print("table font",sz.most_common(4)); print("table linesp",ls.most_common(4)); print("space_after",sa.most_common(3))
c=Counter()
for p in d.paragraphs:
    if p.text.strip(): c[(p.style.name,str(p.paragraph_format.line_spacing),str(p.paragraph_format.space_after))]+=1
print(c.most_common(8))
for n in ("Normal","Body Text","First Paragraph","Table"):
    try:
        s=d.styles[n];print(n,s.font.size,s.font.name,s.paragraph_format.line_spacing,s.paragraph_format.space_after)
    except KeyError: pass
import re
print(re.findall(r'<w:tblCellMar>.*?</w:tblCellMar>',d.tables[5]._tbl.xml,re.S)[:1], d.tables[5].style.name if d.tables[5].style else None)
