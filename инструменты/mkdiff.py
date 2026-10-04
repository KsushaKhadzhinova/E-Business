import openpyxl,os,warnings,json,datetime
warnings.filterwarnings("ignore")
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"; T=r"C:\Users\zheny\AppData\Local\Temp\claude\E-----------------------\89f89955-8ca8-4485-9364-f108500ccfc2\scratchpad"
bro=T+r"\broken"; out=T+r"\diffs"
def enc(v):
    if isinstance(v,(datetime.datetime,datetime.date)):
        d=v if isinstance(v,datetime.datetime) else datetime.datetime(v.year,v.month,v.day)
        return {"d":(d-datetime.datetime(1899,12,30)).total_seconds()/86400}
    return v
tot=0
for root,_,fs in os.walk(base):
    for f in sorted(fs):
        if not f.endswith(".xlsx"): continue
        a=openpyxl.load_workbook(os.path.join(bro,f)); b=openpyxl.load_workbook(os.path.join(root,f))
        ch=[]
        for ws in a.worksheets:
            if ws.title not in b.sheetnames: print("MISSING SHEET",f,ws.title); continue
            w2=b[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    o=w2[c.coordinate]
                    if c.value!=o.value or c.number_format!=o.number_format:
                        ch.append({"s":ws.title,"c":c.coordinate,"v":enc(c.value),"f":c.number_format,"fv":c.value!=o.value})
        for ws in b.worksheets:
            if ws.title not in a.sheetnames: print("EXTRA SHEET",f,ws.title)
        json.dump(ch,open(os.path.join(out,f+".json"),"w",encoding="utf-8"),ensure_ascii=False)
        tot+=len(ch); print(f[:45],len(ch))
print("total",tot)
