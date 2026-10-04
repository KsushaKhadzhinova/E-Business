from docx import Document
import re, os, copy
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
GOALS={
"ЛР1":"Определить границы рынка NotaCode и оценить поисковый спрос по данным Google Trends и Яндекс Вордстат: тренд, сезонность и циклы интереса; сопоставить источники и зафиксировать продуктовые, клиентские, географические и канальные границы рынка.",
"ЛР2-3":"Оценить объем рынка веб-инструментов построения и проверки диаграмм формальных нотаций в Республике Беларусь девятью методиками (поисковый спрос, Similarweb, «сверху вниз» и «снизу вверх», потенциальные клиенты, цифровая воронка, маркетплейсы, сценарная оценка, платежеспособность), сопоставить результаты и сформулировать итоговый диапазон.",
"ЛР4":"Оценить конкурентную среду и барьеры входа на рынок NotaCode по адаптированной модели пяти сил Портера с цифровыми факторами и по данным Similarweb; сформулировать вывод о привлекательности рынка и условиях входа.",
"ЛР6":"Сформулировать ценностное предложение и цифровой товар NotaCode, заполнить бизнес-модель по Канвас, описать бизнес-логику, сценарии, страницы и экраны и разделить разработку на релизы по методикам 6.1–6.5.",
}
def setnum(p,new):
    for r in p.runs:
        if re.match(r'\s*\d',r.text):
            r.text=re.sub(r'^\s*\d+(\.\d+)*',new,r.text,count=1); return
    p.runs[0].text=re.sub(r'^\s*\d+(\.\d+)*',new,p.runs[0].text,count=1)
for lr in ["ЛР1","ЛР2-3","ЛР4","ЛР5","ЛР6","ЛР7"]:
    path=os.path.join(base,lr,f"ОТЧЕТ_{lr}_NotaCode.docx")
    d=Document(path); log=[]
    # title date
    for p in d.paragraphs[:25]:
        if re.match(r'Минск, \w+ 2026',p.text.strip()):
            for r in p.runs: r.text=re.sub(r'(июнь|август) 2026','октябрь 2026',r.text)
            log.append("date")
    cur=0; h1seq=0; first_xod=None
    for p in d.paragraphs:
        s=p.style.name
        if s=="Heading 1":
            m=re.match(r'^\s*(\d+)\s+(.*)',p.text)
            if not m: continue
            title=m.group(2)
            if lr=="ЛР7":
                h1seq+=1; cur=h1seq
            elif title.startswith("ЦЕЛЬ"): cur=1
            elif title.startswith("ХОД"): cur=2; first_xod=p
            else: cur=3
            if int(m.group(1))!=cur: setnum(p,str(cur)); log.append(f"H1 {m.group(1)}->{cur} {title[:25]}")
        elif s in("Heading 2","Heading 3","Heading 4"):
            m=re.match(r'^\s*(\d+)((\.\d+)+)\s',p.text)
            if m and int(m.group(1))!=cur:
                setnum(p,str(cur)+m.group(2)); log.append(f"{s[-1]} {m.group(1)}{m.group(2)}->{cur}{m.group(2)}")
    if lr in GOALS and first_xod is not None:
        h=copy.deepcopy(first_xod._p); first_xod._p.addprevious(h)
        from docx.text.paragraph import Paragraph
        hp=Paragraph(h,first_xod._parent)
        for r in hp.runs[1:]: r._r.getparent().remove(r._r)
        hp.runs[0].text="1 ЦЕЛЬ РАБОТЫ"
        np_=copy.deepcopy(first_xod._p); first_xod._p.addprevious(np_)
        bp=Paragraph(np_,first_xod._parent)
        for r in bp.runs[1:]: r._r.getparent().remove(r._r)
        bp.runs[0].text=GOALS[lr]
        try: bp.style=d.styles["First Paragraph"]
        except KeyError: bp.style=d.styles["Body Text"]
        bp.paragraph_format.page_break_before=False
        bp.paragraph_format.first_line_indent=450215
        bp.alignment=3
        first_xod.paragraph_format.page_break_before=True
        log.append("goal added")
    d.save(path)
    print(lr,len(log),"changes;",[x for x in log if x.startswith(("H1","date","goal"))])
