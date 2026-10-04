from docx import Document
import re
d=Document(r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР7\ОТЧЕТ_ЛР7_NotaCode.docx")
txt=[p.text for p in d.paragraphs]
for tb in d.tables:
    for r in tb.rows:
        for c in r.cells: txt.append(c.text)
full="\n".join(txt)
print(len(full),"chars;",len(d.tables),"tables")
for k in ["197 988","197\xa0988","49 893","414 754","2,32","3,12","3,7","3,83","3,0285","60131","не собрано","модельн","смоделир","сгенерир","ДОСНЯТЬ","проверен","Планируется","1 634","65 375","249","16 590"]:
    print(k.replace("\xa0","[nbsp]"),full.count(k))
for p in d.paragraphs:
    if p.style.name.startswith("Heading 1"): print("H1:",p.text)
