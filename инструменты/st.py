import openpyxl,os,warnings
warnings.filterwarnings("ignore")
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
def P(f): return [os.path.join(r,f) for r,_,fs in os.walk(base) if f in fs][0]
wb=openpyxl.load_workbook(P("6_4_stranicy_ekrany_NotaCode.xlsx"),data_only=True)
for sh in ["Страницы_и_экраны","Матрица_сценарий_интерфейс"]:
    ws=wb[sh]; print("==",sh,ws.max_row,ws.max_column)
    for r in (4,5,6): print(r,[str(c.value)[:18] if c.value is not None else None for c in ws[r][:22]])
wb=openpyxl.load_workbook(P("analiz_konkurentov_Levitt_Kano_NotaCode.xlsx"),data_only=True)
for sh in ["07_Стандарт_рынка","08_Преимущества","04_Кано","03_Левитт"]:
    ws=wb[sh]; print("==",sh,ws.max_row)
    for r in (6,7,8): print(r,[str(c.value)[:16] if c.value is not None else None for c in ws[r][:16]])
