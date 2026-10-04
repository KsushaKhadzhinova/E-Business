from docx import Document
from docx.shared import Pt, Mm
from docx.oxml.ns import qn
import os,sys
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
def compact(path):
    d=Document(path)
    for s in d.sections: s.right_margin=Mm(10)
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    pf=p.paragraph_format
                    pf.line_spacing=1.0; pf.space_before=Pt(0); pf.space_after=Pt(0); pf.first_line_indent=Pt(0)
                    for r in p.runs: r.font.size=Pt(10)
    body=d.element.body; removed=0
    for p in list(body.iterchildren(qn('w:p'))):
        if ''.join(p.itertext()).strip(): continue
        if p.xpath('.//w:drawing|.//w:pict|.//w:br|.//w:sectPr|.//w:fldChar|.//w:bookmarkStart'): continue
        nxt=p.getnext(); prv=p.getprevious()
        if nxt is not None and nxt.tag==qn('w:tbl') and prv is not None and prv.tag==qn('w:tbl'): continue
        body.remove(p); removed+=1
    d.save(path); return removed
for lr in ["ЛР1","ЛР2-3","ЛР4","ЛР5","ЛР6","ЛР7"]:
    print(lr,"removed empty:",compact(os.path.join(base,lr,f"ОТЧЕТ_{lr}_NotaCode.docx")))
