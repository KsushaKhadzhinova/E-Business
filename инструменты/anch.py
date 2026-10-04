import openpyxl,os,re,warnings
from docx import Document
warnings.filterwarnings("ignore")
base=r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"
def P(f): return [os.path.join(r,f) for r,_,fs in os.walk(base) if f in fs][0]
def text(lr):
    d=Document(os.path.join(base,lr,f"ОТЧЕТ_{lr}_NotaCode.docx")); t=[p.text for p in d.paragraphs]
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells: t.append(c.text)
    return re.sub(r"[\s\xa0]+"," "," ".join(t))
def fmt_variants(v):
    out=set()
    if isinstance(v,(int,float)):
        for dec in (0,1,2):
            s=f"{v:,.{dec}f}".replace(",", " ").replace(".",",")
            out.add(s); out.add(s.replace(" ",""))
    return out
checks={
 "ЛР4":[("ocenka_konkurencii_NotaCode.xlsx","07_Барьеры_входа",["F13"]),("ocenka_konkurencii_NotaCode.xlsx","02_Конкуренты_SW",["B37","B38","B39","B40","B41"])],
 "ЛР2-3":[("ПК_потенциальные_клиенты_NotaCode.xlsx","04_Расчет",None),("ПС_поисковый_спрос_NotaCode.xlsx","08_Дашборд",None),("СЦ_сценарная_оценка_Similarweb_NotaCode.xlsx","06_Сверка",None)],
}
for lr,lst in checks.items():
    tx=text(lr)
    for f,sh,cells in lst:
        ws=openpyxl.load_workbook(P(f),data_only=True)[sh]
        vals=[]
        if cells: vals=[(c,ws[c].value) for c in cells]
        else:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value,(int,float)) and abs(c.value)>=2 and c.value!=int(c.value) or (isinstance(c.value,(int,float)) and abs(c.value)>=1000): vals.append((c.coordinate,c.value))
        miss=[]; hit=0
        for c,v in vals[:40]:
            if any(x in tx for x in fmt_variants(v)): hit+=1
            else: miss.append((c,round(v,3) if isinstance(v,float) else v))
        print(lr,f[:25],sh,"found",hit,"of",min(len(vals),40),"missing:",miss[:8])
