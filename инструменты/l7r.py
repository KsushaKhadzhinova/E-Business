from docx import Document
import re
d=Document(r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА\ЛР7\ОТЧЕТ_ЛР7_NotaCode.docx")
texts=[p.text for p in d.paragraphs]
for tb in d.tables:
    for r in tb.rows:
        for c in r.cells: texts.append(c.text)
pat=re.compile(r"(?i)(табл(?:ица|ице|ицы|ицу|\.)?)\s*(\d+)")
own=0; foreign=0; samples=[]
for t in texts:
    if re.match(r"^Таблица \d+ [–-]",t): continue
    for m in pat.finditer(t):
        ctx=t[max(0,m.start()-25):m.start()]
        if re.search(r"ЛР\s*\d|ЛР2-3|методик|шаблон|\[",ctx): foreign+=1
        else:
            own+=1
            if len(samples)<8: samples.append(t[max(0,m.start()-40):m.end()+20])
print("own refs",own,"foreign",foreign); print(samples)
