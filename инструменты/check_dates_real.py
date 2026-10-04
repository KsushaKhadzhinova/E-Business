import openpyxl
import os
from datetime import datetime, date

base = r"E:\ИИТ\ЭлектронныйБизнес\СДАЧА"

print("=== Ввод_данных col A — строки 5-20 (где должны быть даты) ===\n")

for root, dirs, files in os.walk(base):
    for f in sorted(files):
        if not f.endswith('.xlsx'):
            continue
        path = os.path.join(root, f)
        rel = os.path.relpath(path, base)
        try:
            wb = openpyxl.load_workbook(path, data_only=False)
            if 'Ввод_данных' not in wb.sheetnames:
                wb.close()
                continue
            ws = wb['Ввод_данных']
            bad = []
            good = 0
            for row in ws.iter_rows(min_row=5, max_row=20, min_col=1, max_col=1):
                cell = row[0]
                v = cell.value
                if v is None:
                    continue
                t = type(v).__name__
                fmt = cell.number_format
                # Должны быть datetime — если стали строкой с датой, это сломано
                if isinstance(v, str) and (len(v) == 10 and ('.' in v or '-' in v)):
                    bad.append(f"  {cell.coordinate}: СТРОКА={v!r}  ← СЛОМАНО")
                elif isinstance(v, (datetime, date)):
                    good += 1
                elif isinstance(v, int) and v > 40000:
                    bad.append(f"  {cell.coordinate}: serial={v}  fmt={fmt}  ← НУЖНО ПОЧИНИТЬ")
            if bad:
                print(f"ПРОБЛЕМЫ: {rel}")
                for b in bad:
                    print(b)
                print()
            else:
                if good > 0:
                    print(f"OK ({good} datetime-ячеек): {rel}")
            wb.close()
        except Exception as e:
            print(f"ERROR {rel}: {e}")
